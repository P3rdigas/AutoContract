import os
from datetime import datetime

import pandas as pd
from docx import Document
from docx2pdf import convert


def replace_text_in_docx(doc, old_text, new_text):
    for paragraph in doc.paragraphs:
        if old_text in paragraph.text:
            for run in paragraph.runs:
                run.text = run.text.replace(old_text, new_text)


def format_date(date_value, output_format="%d/%m/%Y"):
    # Convert the date to a string with the desired format
    if pd.notna(date_value):
        date_object = pd.to_datetime(date_value, dayfirst=True)
        return date_object.strftime(output_format)
    else:
        return ""


def create_document(
    person_name,
    dob,
    passport,
    pp_ex,
    template_path,
    output_dir,
):
    doc = Document(template_path)

    # Replace placeholders with variable values
    replace_text_in_docx(doc, "CANDIDATE_NAME", person_name)

    # Format the date before replacing it in the document
    formatted_dob = format_date(dob, "%d/%m/%Y")
    replace_text_in_docx(doc, "DOB", formatted_dob)

    replace_text_in_docx(doc, "PASSPORT", passport)

    # Format the date before replacing it in the document
    formatted_pp_ex = format_date(pp_ex, "%d/%m/%Y")
    replace_text_in_docx(doc, "PP_EX", formatted_pp_ex)

    # Output file path
    output_path = os.path.join(output_dir, f"PTC_EN_{person_name}.docx")

    try:
        # Save the modified document
        doc.save(output_path)
        print(f"Document for {person_name} saved to: {output_path}")
        # TODO: convert is slow
        # TODO: Generate filename inside PDF Reader with a lot of 0000\0000
        convert(output_path)
    except PermissionError as e:
        print(f"Permission error: {e}")


def main():
    # Load the Word template
    template_path = os.path.join(
        os.path.dirname(__file__), "input", "template_en_unlock.docx"
    )
    if not os.path.isfile(template_path):
        print(f"Error: Template file not found at {template_path}")
        return

    # Read data from Excel file
    excel_path = os.path.join(os.path.dirname(__file__), "input", "data.xlsx")
    if not os.path.isfile(excel_path):
        print(f"Error: Excel file not found at {excel_path}")
        return

    # Specify the column names
    column_names = [
        "SL",
        "CANDIDATE NAME",
        "POSITION",
        "PASSPORT",
        "PP EX",
        "DOB",
        "REMARKS",
    ]

    # Read data from Excel file using specified column names
    df = pd.read_excel(excel_path, usecols=column_names)

    # Output directory
    output_dir = os.path.join(os.path.dirname(__file__), "output")
    os.makedirs(output_dir, exist_ok=True)

    # Iterate through rows and create Word documents
    for index, row in df.iterrows():
        person_name = row["CANDIDATE NAME"]
        dob = row["DOB"]
        passport = str(row["PASSPORT"])
        passport_exp = row["PP EX"]

        create_document(
            person_name, dob, passport, passport_exp, template_path, output_dir
        )


if __name__ == "__main__":
    main()
